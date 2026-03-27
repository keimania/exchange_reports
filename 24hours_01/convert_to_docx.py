#!/usr/bin/env python3
"""
MD -> DOCX Report Converter (v2)
- Clickable hyperlinks for URLs
- .md references converted to .docx paths in appendix
- Professional formatting with cover page
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPORTS_DIR = Path(__file__).parent / "reports"
OUTPUT_DIR = Path(__file__).parent / "output" / "docx"

REPORT_FILES = [
    ("00_종합비교_요약.md", "종합 비교 요약"),
    ("01_KRX_한국거래소.md", "KRX 한국거래소"),
    ("02_NYSE_뉴욕증권거래소.md", "NYSE 뉴욕증권거래소"),
    ("03_NASDAQ_나스닥.md", "NASDAQ 나스닥"),
    ("04_24X_Exchange.md", "24X Exchange"),
    ("05_CME_Group.md", "CME Group"),
    ("06_DTCC_청산결제.md", "DTCC 청산결제"),
    ("07_LSE_런던증권거래소.md", "LSE 런던증권거래소"),
    ("08_JPX_도쿄증권거래소.md", "JPX 도쿄증권거래소"),
    ("09_Blue_Ocean_ATS.md", "Blue Ocean ATS"),
    ("10_유럽거래소_Eurex_Euronext.md", "유럽거래소 Eurex/Euronext"),
]

ACCENT = RGBColor(0, 102, 204)
DARK = RGBColor(26, 26, 46)
GRAY = RGBColor(100, 100, 100)
LINK_COLOR = RGBColor(0, 102, 204)


def add_hyperlink(paragraph, text, url):
    """Add a clickable hyperlink to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color')
    c.set(qn('w:val'), '0066CC')
    rPr.append(c)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '18')  # 9pt
    rPr.append(sz)
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Malgun Gothic')
    rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    rPr.append(rFonts)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._element.append(hyperlink)


def setup_styles(doc):
    """Set up document styles."""
    style = doc.styles['Normal']
    style.font.name = 'Malgun Gothic'
    style.font.size = Pt(10)
    style.font.color.rgb = DARK
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(4)
    rpr = style.element.get_or_add_rPr()
    ea = rpr.makeelement(qn('w:rFonts'), {})
    ea.set(qn('w:eastAsia'), 'Malgun Gothic')
    rpr.append(ea)

    for level, (size, color) in {
        'Heading 1': (22, ACCENT),
        'Heading 2': (16, ACCENT),
        'Heading 3': (13, DARK),
        'Heading 4': (11, DARK),
    }.items():
        s = doc.styles[level]
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = color
        s.font.name = 'Malgun Gothic'
        s.paragraph_format.space_before = Pt(14)
        s.paragraph_format.space_after = Pt(6)


def add_cover_page(doc, title):
    """Add a cover page."""
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("24시간 거래 체제 비교 분석 리포트")
    run.font.size = Pt(28)
    run.font.color.rgb = ACCENT
    run.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.size = Pt(18)
    run.font.color.rgb = DARK

    for _ in range(3):
        doc.add_paragraph()

    for text in ["KRX vs 글로벌 주요 거래소", "2026년 3월 기준 | EXTURE 3.0 기반 분석"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(11 if "KRX" in text else 10)
        run.font.color.rgb = GRAY

    doc.add_page_break()


def parse_table(lines):
    """Parse markdown table lines into rows of cells."""
    rows = []
    for line in lines:
        line = line.strip()
        if line.startswith('|') and line.endswith('|'):
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if all(re.match(r'^[-:]+$', c) for c in cells):
                continue
            rows.append(cells)
    return rows


def add_table_to_doc(doc, rows):
    """Add a formatted table."""
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            if j >= num_cols:
                continue
            cell = table.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            clean = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_text)
            clean = re.sub(r'\*(.*?)\*', r'\1', clean)

            # Check for markdown links in cell
            link_match = re.search(r'\[([^\]]+)\]\(([^)]+)\)', clean)
            if link_match:
                before = clean[:link_match.start()]
                link_text = link_match.group(1)
                link_url = link_match.group(2)
                after = clean[link_match.end():]
                if before:
                    run = p.add_run(before)
                    run.font.size = Pt(8.5)
                    run.font.name = 'Malgun Gothic'
                if link_url.startswith('http'):
                    add_hyperlink(p, link_text, link_url)
                elif link_url.endswith('.md'):
                    docx_path = link_url.replace('.md', '.docx')
                    run = p.add_run(f"{link_text} ({docx_path})")
                    run.font.size = Pt(8.5)
                    run.font.name = 'Malgun Gothic'
                    run.font.color.rgb = LINK_COLOR
                else:
                    run = p.add_run(link_text)
                    run.font.size = Pt(8.5)
                    run.font.name = 'Malgun Gothic'
                if after:
                    run = p.add_run(after)
                    run.font.size = Pt(8.5)
                    run.font.name = 'Malgun Gothic'
            else:
                run = p.add_run(clean)
                run.font.size = Pt(8.5)
                run.font.name = 'Malgun Gothic'

            if i == 0:
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                shading = cell._element.get_or_add_tcPr()
                bg = shading.makeelement(qn('w:shd'), {qn('w:fill'): '0066CC', qn('w:val'): 'clear'})
                shading.append(bg)
            elif i % 2 == 0:
                shading = cell._element.get_or_add_tcPr()
                bg = shading.makeelement(qn('w:shd'), {qn('w:fill'): 'F0F4F8', qn('w:val'): 'clear'})
                shading.append(bg)

    doc.add_paragraph()


def add_text_with_links(paragraph, text):
    """Add text to paragraph, converting markdown links and bold to runs."""
    # Split by markdown link pattern
    parts = re.split(r'(\[.*?\]\(.*?\))', text)
    for part in parts:
        link_match = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
        if link_match:
            link_text = link_match.group(1)
            url = link_match.group(2)
            if url.startswith('http'):
                add_hyperlink(paragraph, link_text, url)
            elif url.endswith('.md'):
                docx_name = url.replace('.md', '.docx')
                run = paragraph.add_run(f"{link_text} ({docx_name})")
                run.font.color.rgb = LINK_COLOR
            else:
                paragraph.add_run(link_text)
        else:
            # Handle bold
            bold_parts = re.split(r'(\*\*.*?\*\*)', part)
            for bp in bold_parts:
                if bp.startswith('**') and bp.endswith('**'):
                    run = paragraph.add_run(bp[2:-2])
                    run.font.bold = True
                    run.font.color.rgb = ACCENT
                else:
                    paragraph.add_run(bp)


def md_to_docx(md_content, title):
    """Convert markdown content to a DOCX document."""
    doc = Document()
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)

    setup_styles(doc)
    add_cover_page(doc, title)

    lines = md_content.split('\n')
    i = 0
    table_buffer = []
    in_table = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('|') and stripped.endswith('|'):
            table_buffer.append(stripped)
            in_table = True
            i += 1
            continue
        elif in_table:
            rows = parse_table(table_buffer)
            if rows:
                add_table_to_doc(doc, rows)
            table_buffer = []
            in_table = False

        if stripped in ('---', '***', '___'):
            i += 1
            continue

        if stripped.startswith('####'):
            doc.add_heading(stripped.lstrip('#').strip(), level=4)
        elif stripped.startswith('###'):
            doc.add_heading(stripped.lstrip('#').strip(), level=3)
        elif stripped.startswith('##'):
            doc.add_heading(stripped.lstrip('#').strip(), level=2)
        elif stripped.startswith('#'):
            doc.add_heading(stripped.lstrip('#').strip(), level=1)
        elif stripped.startswith('>'):
            text = stripped.lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run(text)
            run.font.italic = True
            run.font.color.rgb = GRAY
        elif stripped.startswith('- ') or stripped.startswith('* '):
            text = stripped[2:]
            p = doc.add_paragraph(style='List Bullet')
            add_text_with_links(p, text)
        elif re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped)
            p = doc.add_paragraph(style='List Number')
            add_text_with_links(p, text)
        elif stripped:
            p = doc.add_paragraph()
            add_text_with_links(p, stripped)

        i += 1

    if table_buffer:
        rows = parse_table(table_buffer)
        if rows:
            add_table_to_doc(doc, rows)

    return doc


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    for filename, title in REPORT_FILES:
        md_path = REPORTS_DIR / filename
        if not md_path.exists():
            print(f"  [SKIP] {filename}")
            continue
        md_content = md_path.read_text(encoding='utf-8')
        doc = md_to_docx(md_content, title)
        out_path = OUTPUT_DIR / filename.replace('.md', '.docx')
        doc.save(str(out_path))
        print(f"  [OK] {out_path.name}")

    print(f"\n  Output: {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
