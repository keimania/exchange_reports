#!/usr/bin/env python3
"""
build_reports.py v4
  - Body [REF:#N] → <sup>[N]</sup> linking to #ref-N in bibliography
  - Bibliography section only: anchor ids + external URL links
  - Sidebar TOC: no citation tags
  - Index: Korean first
"""

import re, sys, markdown
from pathlib import Path
from datetime import datetime

BASE_DIR  = Path(__file__).resolve().parent
MD_DIR    = BASE_DIR / "deliverables" / "md"
HTML_DIR  = BASE_DIR / "deliverables" / "html"
DOCX_DIR  = BASE_DIR / "deliverables" / "docx"
GLOSS_SRC = BASE_DIR / "glossary.md"

HTML_DIR.mkdir(parents=True, exist_ok=True)
DOCX_DIR.mkdir(parents=True, exist_ok=True)

# ── CSS ───────────────────────────────────────────────────────────────────
CSS = r"""
:root{--bg:#fff;--fg:#1a1a2e;--accent:#0f3460;--border:#ddd;
  --th:#0f3460;--thf:#fff;--stripe:#f5f7fa;--link:#1565c0;--code:#f4f4f4;--side:260px}
*{box-sizing:border-box;margin:0;padding:0}
html{font-size:15px;scroll-behavior:smooth}
body{font-family:"Pretendard","Noto Sans KR",-apple-system,BlinkMacSystemFont,
  "Segoe UI",Roboto,sans-serif;color:var(--fg);background:var(--bg);line-height:1.72}
.page-wrap{display:flex;min-height:100vh}
aside#sidebar{
  width:var(--side);min-width:var(--side);position:sticky;top:0;height:100vh;
  overflow-y:auto;border-right:1px solid var(--border);padding:1.2rem .8rem;
  background:#fafbfc;font-size:.82rem;flex-shrink:0;
}
aside#sidebar h3{font-size:.9rem;color:var(--accent);margin:0 0 .6rem;
  padding-bottom:.3rem;border-bottom:2px solid var(--accent)}
aside#sidebar ul{list-style:none;padding:0}
aside#sidebar li{margin:.15rem 0}
aside#sidebar li.h3{padding-left:.8rem;font-size:.78rem}
aside#sidebar a{color:#333;text-decoration:none;display:block;padding:2px 4px;border-radius:3px}
aside#sidebar a:hover,aside#sidebar a.active{background:var(--accent);color:#fff}
aside#sidebar .home-link{display:block;margin-top:1rem;padding-top:.6rem;
  border-top:1px solid var(--border);font-weight:600;color:var(--accent)}
main#content{flex:1;max-width:780px;padding:2rem 2.5rem 4rem}
h1{font-size:1.85rem;border-bottom:3px solid var(--accent);padding-bottom:.4rem;margin:1.5rem 0 1rem}
h2{font-size:1.4rem;color:var(--accent);margin:1.6rem 0 .5rem;border-bottom:1px solid var(--border);padding-bottom:.2rem}
h3{font-size:1.12rem;margin:1.1rem 0 .35rem}
h4{font-size:1rem;margin:.9rem 0 .25rem}
p{margin:.45rem 0}
a{color:var(--link);text-decoration:none}a:hover{text-decoration:underline}
hr{border:none;border-top:1px solid var(--border);margin:1.4rem 0}
code{background:var(--code);padding:2px 5px;border-radius:3px;font-size:.9em}
pre{background:var(--code);padding:1rem;border-radius:6px;overflow-x:auto;margin:.7rem 0}
pre code{background:none;padding:0}
blockquote{border-left:4px solid var(--accent);padding:.4rem 1rem;margin:.7rem 0;color:#555;background:#f9f9fb}
ul,ol{margin:.35rem 0 .35rem 1.5rem}li{margin:.12rem 0}
strong{color:var(--accent)}
table{width:100%;border-collapse:collapse;margin:.7rem 0;font-size:.86rem}
th{background:var(--th);color:var(--thf);padding:.45rem .55rem;text-align:left;font-weight:600}
td{padding:.4rem .55rem;border-bottom:1px solid var(--border)}
tr:nth-child(even) td{background:var(--stripe)}
sup.cite{font-size:.72em;vertical-align:super;margin-left:1px}
sup.cite a{color:var(--accent);font-weight:600;text-decoration:none;
  border-bottom:1px dotted var(--accent)}
sup.cite a:hover{color:var(--link);border-bottom-style:solid}
a[target="_blank"]::after{content:"";font-size:.7em;margin-left:2px}
a[target="_blank"]{color:#666;word-break:break-all;font-size:.85em}
footer{margin-top:2rem;padding-top:.8rem;border-top:1px solid var(--border);font-size:.78rem;color:#999}
@media(max-width:860px){aside#sidebar{display:none}main#content{max-width:100%;padding:1.2rem}}
@media print{aside#sidebar{display:none}main#content{max-width:100%;padding:1cm}
  sup.cite a{color:#000;border:none}a{color:var(--fg);text-decoration:none}}
"""

# ── Helpers ───────────────────────────────────────────────────────────────
def slugify(t):
    t = re.sub(r'[^\w\s가-힣-]', '', t.lower())
    return re.sub(r'\s+', '-', t.strip())[:80]

def inject_ids(html):
    ctr = {}
    def rpl(m):
        tag, inner = m.group(1), m.group(2)
        clean = re.sub(r'<[^>]+>', '', inner)
        s = slugify(clean) or "section"
        ctr[s] = ctr.get(s, 0) + 1
        if ctr[s] > 1: s = f"{s}-{ctr[s]}"
        return f'<{tag} id="{s}">{inner}</{tag}>'
    return re.sub(r'<(h[1234])>(.*?)</\1>', rpl, html)

def strip_cite_tags(text):
    text = re.sub(r'<sup class="cite">.*?</sup>', '', text)
    text = re.sub(r'\[REF:#?\d+[^\]]*\]', '', text)
    text = re.sub(r'\[WEB:[^\]]*\]', '', text)
    return text.strip()

def build_sidebar(html):
    hdgs = re.findall(r'<(h[23])\b[^>]*id="([^"]*)"[^>]*>(.*?)</\1>', html)
    if not hdgs: return ""
    items = []
    for tag, anchor, raw in hdgs:
        clean = strip_cite_tags(re.sub(r'<[^>]+>', '', raw))
        if not clean: continue
        cls = ' class="h3"' if tag == "h3" else ""
        items.append(f'<li{cls}><a href="#{anchor}">{clean}</a></li>')
    return '\n'.join(items)

# ── Bibliography parsing ──────────────────────────────────────────────────
def parse_bibliography(md_text):
    ref_map = {}
    web_map = {}
    for m in re.finditer(r'\[REF:#(\d+)\]\s+(.+?)(?:\n|$)', md_text):
        u = re.search(r'(https?://\S+)', m.group(2))
        ref_map[m.group(1)] = u.group(1).rstrip(',.') if u else None
    for m in re.finditer(r'\[WEB:([^\]]+)\]\s+(.+?)(?:\n|$)', md_text):
        u = re.search(r'(https?://\S+)', m.group(2))
        web_map[m.group(1)] = u.group(1).rstrip(',.') if u else None
    return ref_map, web_map

def web_anchor(key):
    return re.sub(r'[^\w-]', '-', key.lower()).strip('-')

# ── Two-pass citation processing ──────────────────────────────────────────

def find_bib_start(md_text):
    """Find the line index where bibliography section starts in markdown."""
    patterns = [
        r'^##\s+10\.\s+Source Bibliography',
        r'^##\s+10\.\s+출처',
        r'^###\s+Reference Sources',
        r'^###\s+참조 출처',
        r'^##\s+10\.\s+출처 목록',
    ]
    for i, line in enumerate(md_text.split('\n')):
        for pat in patterns:
            if re.match(pat, line.strip()):
                return i
    return None

def process_md_two_pass(md_text):
    """
    Split md into body + bibliography.
    Convert each separately so bibliography anchors are unique.
    """
    bib_start = find_bib_start(md_text)
    lines = md_text.split('\n')

    if bib_start is not None:
        body_md = '\n'.join(lines[:bib_start])
        bib_md = '\n'.join(lines[bib_start:])
    else:
        body_md = md_text
        bib_md = ""

    return body_md, bib_md

def render_body(body_md):
    """Render body: convert [REF:#N] to <sup> linking #ref-N, [WEB:key] to #web-key."""
    exts = ['tables', 'fenced_code', 'toc', 'sane_lists']
    html = markdown.markdown(body_md, extensions=exts)
    html = inject_ids(html)

    # [REF:#N ...] → <sup><a href="#ref-N">[N]</a></sup>
    def ref_rpl(m):
        num = m.group(1)
        extra = (m.group(2) or "").strip()
        tip = f"REF #{num}" + (f" {extra}" if extra else "")
        return f'<sup class="cite"><a href="#ref-{num}" title="{tip}">[{num}]</a></sup>'
    html = re.sub(r'\[REF:#(\d+)([^\]]*)\]', ref_rpl, html)

    # [WEB:key] → <sup><a href="#web-key">[W]</a></sup>
    def web_rpl(m):
        key = m.group(1)
        anc = web_anchor(key)
        return f'<sup class="cite"><a href="#web-{anc}" title="WEB: {key}">[W]</a></sup>'
    html = re.sub(r'\[WEB:([^\]]+)\]', web_rpl, html)

    return html

def render_bibliography(bib_md, ref_map, web_map):
    """
    Render bibliography section:
    - Each [REF:#N] gets a unique id="ref-N" anchor
    - Each [WEB:key] gets a unique id="web-key" anchor
    - URLs become clickable external links
    No duplicate ids because this only runs on the bibliography text.
    """
    exts = ['tables', 'fenced_code', 'toc', 'sane_lists']
    html = markdown.markdown(bib_md, extensions=exts)
    html = inject_ids(html)

    # Process REF entries: add anchor id, make URLs clickable
    used_refs = set()
    def ref_bib(m):
        num = m.group(1)
        rest = m.group(2)
        # Only add id for first occurrence of each ref number
        if num not in used_refs:
            used_refs.add(num)
            anchor = f' id="ref-{num}"'
        else:
            anchor = ""
        # Make URLs clickable
        rest = re.sub(
            r'(https?://\S+)',
            lambda u: f'<a href="{u.group(1).rstrip(",.")}" target="_blank" rel="noopener noreferrer">{u.group(1).rstrip(",.")} &#8599;</a>',
            rest
        )
        return f'<span{anchor}><strong>[REF:#{num}]</strong></span> {rest}'

    html = re.sub(r'\[REF:#(\d+)\]\s*(.+?)(?=<br|</p|</li|\n|$)', ref_bib, html)

    # Process WEB entries: add anchor id, make URLs clickable
    used_webs = set()
    def web_bib(m):
        key = m.group(1)
        rest = m.group(2)
        anc = web_anchor(key)
        if anc not in used_webs:
            used_webs.add(anc)
            anchor = f' id="web-{anc}"'
        else:
            anchor = ""
        rest = re.sub(
            r'(https?://\S+)',
            lambda u: f'<a href="{u.group(1).rstrip(",.")}" target="_blank" rel="noopener noreferrer">{u.group(1).rstrip(",.")} &#8599;</a>',
            rest
        )
        return f'<span{anchor}><strong>[WEB:{key}]</strong></span> {rest}'

    html = re.sub(r'\[WEB:([^\]]+)\]\s*(.+?)(?=<br|</p|</li|\n|$)', web_bib, html)

    return html


# ── HTML build ────────────────────────────────────────────────────────────
def md_to_full_html(md_text, title, back_link="index.html"):
    ref_map, web_map = parse_bibliography(md_text)
    body_md, bib_md = process_md_two_pass(md_text)

    body_html = render_body(body_md)
    bib_html = render_bibliography(bib_md, ref_map, web_map) if bib_md else ""

    full_body = body_html + "\n" + bib_html
    sidebar_items = build_sidebar(full_body)

    return f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title>
<style>{CSS}</style>
</head>
<body>
<div class="page-wrap">
<aside id="sidebar">
  <h3>Contents</h3>
  <ul>{sidebar_items}</ul>
  <a class="home-link" href="{back_link}">&#8592; Index</a>
</aside>
<main id="content">
{full_body}
<footer>
  Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} &middot;
  Project compare_exchange_01 &middot;
  <a href="{back_link}">Back to Index</a>
</footer>
</main>
</div>
</body>
</html>"""


def collect_md_files():
    files = sorted(MD_DIR.glob("*.md"))
    if GLOSS_SRC.exists():
        files.append(GLOSS_SRC)
    return files

def build_html():
    files = collect_md_files()
    results = []
    for md_path in files:
        md_text = md_path.read_text(encoding="utf-8")
        tm = re.search(r'^#\s+(.+)', md_text, re.MULTILINE)
        title = tm.group(1) if tm else md_path.stem
        if md_path == GLOSS_SRC:
            out_name = "glossary.html"; back = "../index.html"
        else:
            out_name = md_path.stem + ".html"; back = "../index.html"
        full = md_to_full_html(md_text, title, back)
        out_path = HTML_DIR / out_name
        out_path.write_text(full, encoding="utf-8")
        results.append((md_path.name, out_name, len(full)))
    return results


# ── DOCX ──────────────────────────────────────────────────────────────────
def build_docx():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    files = collect_md_files()
    results = []
    for md_path in files:
        md_text = md_path.read_text(encoding="utf-8")
        md_text = re.sub(r'\[REF:#\d+[^\]]*\]', '', md_text)
        md_text = re.sub(r'\[WEB:[^\]]+\]', '', md_text)
        doc = Document()
        style = doc.styles['Normal']
        style.font.name = 'Malgun Gothic'; style.font.size = Pt(10)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
        for lv in range(1, 5):
            hs = doc.styles[f'Heading {lv}']
            hs.font.name = 'Malgun Gothic'
            hs.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
            hs.font.color.rgb = RGBColor(0x0f, 0x34, 0x60)
        lines = md_text.split('\n')
        in_table = False; table_rows = []; in_code = False
        for line in lines:
            s = line.strip()
            if s.startswith('```'): in_code = not in_code; continue
            if in_code: doc.add_paragraph(line); continue
            if in_table and (not s or not s.startswith('|')):
                _flush_table(doc, table_rows); table_rows = []; in_table = False
            if s.startswith('|') and '|' in s[1:]:
                cells = [c.strip() for c in s.strip('|').split('|')]
                if all(set(c) <= set('-: ') for c in cells): continue
                table_rows.append(cells); in_table = True; continue
            hm = re.match(r'^(#{1,4})\s+(.+)', s)
            if hm: doc.add_heading(hm.group(2), level=len(hm.group(1))); continue
            if s in ('---','***','___'): doc.add_paragraph('─'*60); continue
            if s.startswith('>'):
                p = doc.add_paragraph(s.lstrip('> '))
                p.paragraph_format.left_indent = Inches(0.5); continue
            if s.startswith('- ') or s.startswith('* '):
                doc.add_paragraph(s[2:], style='List Bullet'); continue
            if s:
                clean = re.sub(r'\*\*(.+?)\*\*', r'\1', s)
                clean = re.sub(r'`(.+?)`', r'\1', clean)
                clean = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', clean)
                doc.add_paragraph(clean)
        if in_table and table_rows: _flush_table(doc, table_rows)
        stem = "glossary" if md_path == GLOSS_SRC else md_path.stem
        out_path = DOCX_DIR / (stem + ".docx")
        doc.save(str(out_path)); results.append((md_path.name, out_path.name))
    return results

def _flush_table(doc, rows):
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn; from docx.oxml import OxmlElement
    if not rows: return
    nc = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=nc); tbl.style = 'Table Grid'
    for i, rd in enumerate(rows):
        for j, ct in enumerate(rd):
            if j >= nc: continue
            cell = tbl.rows[i].cells[j]; cell.text = ct
            if i == 0:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True; run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'),'0F3460')
                cell._tc.get_or_add_tcPr().append(sh)
            else:
                for p in cell.paragraphs:
                    for run in p.runs: run.font.size = Pt(9)


# ── Index (Korean first) ─────────────────────────────────────────────────
def build_index():
    html_files = sorted(HTML_DIR.glob("*.html"))
    synthesis = []; en = []; ko = []; glossary = None
    for f in html_files:
        if f.name == "index.html": continue
        if f.name == "glossary.html": glossary = f; continue
        if "종합비교" in f.name: synthesis.append(f)
        elif "_ko" in f.name: ko.append(f)
        else: en.append(f)

    def row(hf):
        stem = hf.stem
        dx = DOCX_DIR / (stem + ".docx")
        dl = f'<a href="docx/{dx.name}">DOCX</a>' if dx.exists() else '—'
        return f'<tr><td><a href="html/{hf.name}">{stem}</a></td><td>{dl}</td></tr>'

    gloss_row = ""
    if glossary:
        dx = DOCX_DIR / "glossary.docx"
        dl = f'<a href="docx/glossary.docx">DOCX</a>' if dx.exists() else '—'
        gloss_row = f'<tr><td><a href="html/glossary.html">Glossary / 용어집 (140 terms)</a></td><td>{dl}</td></tr>'

    page = f"""<!DOCTYPE html>
<html lang="ko"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>9대 거래소 시장 미시구조 비교</title>
<style>{CSS}
aside#sidebar{{display:none}}
main#content{{max-width:960px;margin:0 auto;padding:2rem 1.5rem}}
.cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:.8rem;margin:1.2rem 0}}
.card{{background:#fafbfc;border:1px solid var(--border);border-radius:8px;padding:.9rem;text-align:center}}
.card .n{{font-size:1.9rem;font-weight:700;color:var(--accent)}}
.card .l{{font-size:.82rem;color:#666}}
</style></head><body><div class="page-wrap"><main id="content">
<h1>9대 거래소 시장 미시구조 비교</h1>
<p style="color:#555">KRX &middot; NYSE &middot; NASDAQ &middot; TSE &middot; LSE &middot; HKEX &middot; Eurex &middot; ASX &middot; SGX</p>
<div class="cards">
<div class="card"><div class="n">9</div><div class="l">거래소</div></div>
<div class="card"><div class="n">140</div><div class="l">용어</div></div>
<div class="card"><div class="n">~197</div><div class="l">출처</div></div>
<div class="card"><div class="n">{len(html_files)}</div><div class="l">보고서</div></div>
</div>
<h2>종합비교</h2>
<table><tr><th>보고서</th><th>DOCX</th></tr>
{''.join(row(f) for f in synthesis)}</table>
<h2>용어집</h2>
<table><tr><th>보고서</th><th>DOCX</th></tr>{gloss_row}</table>
<h2>한국어 보고서</h2>
<table><tr><th>보고서</th><th>DOCX</th></tr>
{''.join(row(f) for f in ko)}</table>
<h2>English Reports</h2>
<table><tr><th>Report</th><th>DOCX</th></tr>
{''.join(row(f) for f in en)}</table>
<footer>Project compare_exchange_01 &middot; Generated {datetime.now().strftime('%Y-%m-%d')}</footer>
</main></div></body></html>"""

    out = BASE_DIR / "deliverables" / "index.html"
    out.write_text(page, encoding="utf-8")
    return str(out)


# ── QA ────────────────────────────────────────────────────────────────────
def run_qa():
    R = []
    md_files = collect_md_files()

    # 1 HTML exist
    ok = True
    for f in md_files:
        stem = "glossary" if f == GLOSS_SRC else f.stem
        h = HTML_DIR / (stem + ".html")
        if not h.exists() or h.stat().st_size < 100:
            R.append(("FAIL", f"HTML missing: {h.name}")); ok = False
    if ok: R.append(("PASS", f"All {len(md_files)} HTML files exist"))

    # 2 DOCX exist
    ok = True
    for f in md_files:
        stem = "glossary" if f == GLOSS_SRC else f.stem
        d = DOCX_DIR / (stem + ".docx")
        if not d.exists() or d.stat().st_size < 1000:
            R.append(("FAIL", f"DOCX missing: {d.name}")); ok = False
    if ok: R.append(("PASS", f"All {len(md_files)} DOCX files exist"))

    # 3 Index Korean first
    idx = BASE_DIR / "deliverables" / "index.html"
    if idx.exists():
        t = idx.read_text("utf-8")
        ko_pos = t.find('한국어 보고서')
        en_pos = t.find('English Reports')
        if ko_pos > 0 and en_pos > 0 and ko_pos < en_pos:
            R.append(("PASS", "index.html: Korean before English"))
        else:
            R.append(("FAIL", "index.html: ordering wrong"))
    else:
        R.append(("FAIL", "index.html missing"))

    # 4 No duplicate ref-* ids in any HTML
    dup_total = 0
    for h in HTML_DIR.glob("*.html"):
        t = h.read_text("utf-8")
        from collections import Counter
        ids = re.findall(r'id="(ref-\d+|web-[^"]+)"', t)
        dupes = {k: v for k, v in Counter(ids).items() if v > 1}
        if dupes:
            dup_total += sum(v - 1 for v in dupes.values())
    if dup_total == 0:
        R.append(("PASS", "No duplicate ref/web anchor IDs"))
    else:
        R.append(("FAIL", f"{dup_total} duplicate anchor IDs found"))

    # 5 Body citations point to bibliography (not external, not to other body cites)
    sample = HTML_DIR / "01_KRX.html"
    if sample.exists():
        t = sample.read_text("utf-8")
        cites = re.findall(r'<sup class="cite"><a href="([^"]+)"', t)
        ext = [c for c in cites if c.startswith('http')]
        internal = [c for c in cites if c.startswith('#')]
        if ext:
            R.append(("FAIL", f"{len(ext)} body cites link externally"))
        elif internal:
            # Verify each internal link target exists as an id in the bib section
            all_ids = set(re.findall(r'id="([^"]+)"', t))
            broken = [c for c in internal if c[1:] not in all_ids]
            if broken:
                R.append(("FAIL", f"{len(broken)} broken internal links"))
            else:
                R.append(("PASS", f"All {len(internal)} body cites → valid bib anchors"))
        else:
            R.append(("PASS", "No citations (report may lack refs)"))

    # 6 Bibliography has external links
    if sample.exists():
        t = sample.read_text("utf-8")
        bib_ext = len(re.findall(r'<a href="https?://[^"]*" target="_blank"', t))
        if bib_ext > 0:
            R.append(("PASS", f"Bibliography: {bib_ext} external links"))
        else:
            R.append(("FAIL", "No external links in bibliography"))

    # 7 Sidebar clean
    if sample.exists():
        t = sample.read_text("utf-8")
        sb = re.search(r'<aside id="sidebar">(.*?)</aside>', t, re.S)
        if sb:
            sbt = sb.group(1)
            if '[REF' in sbt or '[WEB' in sbt or 'class="cite"' in sbt:
                R.append(("FAIL", "Sidebar has citation tags"))
            else:
                R.append(("PASS", "Sidebar clean"))

    # 8 Glossary
    g = HTML_DIR / "glossary.html"
    if g.exists() and g.stat().st_size > 5000:
        R.append(("PASS", "glossary.html included"))
    else:
        R.append(("FAIL", "glossary.html missing"))

    # 9 No CDN
    ok = True
    for h in HTML_DIR.glob("*.html"):
        t = h.read_text("utf-8")
        if 'cdn.' in t.lower() or 'googleapis.com' in t:
            R.append(("FAIL", f"CDN in {h.name}")); ok = False
    if ok: R.append(("PASS", "Fully offline"))

    return R


def main():
    args = set(sys.argv[1:])
    do_all = not args or '--all' in args
    if do_all or '--html' in args:
        print("=== HTML ===")
        for n, o, s in build_html(): print(f"  {n:32s} -> {o:32s} ({s:,}b)")
        print()
    if do_all or '--docx' in args:
        print("=== DOCX ===")
        for n, o in build_docx(): print(f"  {n:32s} -> {o}")
        print()
    if do_all or '--index' in args:
        print("=== Index ===")
        print(f"  {build_index()}\n")
    if do_all or '--qa' in args:
        print("=== QA ===")
        qa = run_qa()
        p = sum(1 for s, _ in qa if s == "PASS")
        f = sum(1 for s, _ in qa if s == "FAIL")
        for s, m in qa: print(f"  {'✅' if s=='PASS' else '❌'} [{s}] {m}")
        print(f"\n  {p} passed, {f} failed")
        if f: sys.exit(1)
    print("\n=== Done ===")

if __name__ == "__main__":
    main()
